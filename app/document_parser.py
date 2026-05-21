"""
多格式文档解析器 —— SOP 极速版
支持 PDF / Word / Excel（含 .xls）/ PowerPoint / TXT / 图片 OCR
"""
import io
import logging
import os
from dataclasses import dataclass
from typing import List

import fitz  # PyMuPDF
from PIL import Image

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    file_name: str
    total_pages: int
    full_text: str
    table_headers: dict = None  # {"Sheet1": "工号 | 姓名 | 基本工资 | ..."}

    def __post_init__(self):
        if self.table_headers is None:
            self.table_headers = {}


def parse_document(file_bytes: bytes, file_name: str) -> ParsedDocument:
    ext = os.path.splitext(file_name)[1].lower()
    logger.info("解析文档: %s (%.1f KB)", file_name, len(file_bytes) / 1024)

    if ext == ".pdf":
        return _parse_pdf(file_bytes, file_name)
    elif ext in (".docx", ".doc"):
        return _parse_docx(file_bytes, file_name)
    elif ext == ".xlsx":
        return _parse_xlsx(file_bytes, file_name)
    elif ext == ".xls":
        return _parse_xls(file_bytes, file_name)
    elif ext in (".pptx", ".ppt"):
        return _parse_pptx(file_bytes, file_name)
    elif ext in (".txt", ".md", ".csv", ".log", ".json", ".xml", ".html", ".htm"):
        return _parse_text(file_bytes, file_name)
    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"):
        return _parse_image(file_bytes, file_name)
    else:
        return _parse_text(file_bytes, file_name)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _parse_pdf(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        doc: fitz.Document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ValueError(f"无法解析 PDF: {file_name}") from exc

    if doc.page_count == 0:
        doc.close()
        raise ValueError(f"PDF 无页面: {file_name}")

    parts = []
    for i in range(doc.page_count):
        try:
            page = doc.load_page(i)
            text = page.get_text("text")
            if text.strip():
                parts.append(f"[第{i+1}页] {text.strip()}")
        except Exception:
            continue

    doc.close()
    full_text = "\n\n".join(parts)
    return ParsedDocument(file_name=file_name, total_pages=len(parts), full_text=full_text)


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def _parse_docx(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        full_text = "\n\n".join(paragraphs)
        return ParsedDocument(file_name=file_name, total_pages=1, full_text=full_text)
    except Exception as exc:
        raise ValueError(f"无法解析 Word 文档: {file_name}") from exc


# ---------------------------------------------------------------------------
# Excel (.xlsx)
# ---------------------------------------------------------------------------

def _parse_xlsx(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        parts = []
        headers = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            # 找表头行：第一个非空单元格 >= 2 的行
            header_row = None
            header_cells = []
            for i, row in enumerate(rows):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(cells) >= 2:
                    header_row = i
                    header_cells = cells
                    break
            if header_row is not None:
                header_text = " | ".join(header_cells)
                headers[sheet_name] = header_text
            # 构建行文本
            rows_text = []
            for i, row in enumerate(rows):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if not cells:
                    continue
                if i == header_row:
                    # 表头行保留原样
                    rows_text.append(" | ".join(cells))
                elif header_cells and len(cells) <= len(header_cells):
                    # 数据行：转为 "表头: 值" 格式
                    pairs = []
                    for j, val in enumerate(cells):
                        h = header_cells[j] if j < len(header_cells) else f"列{j+1}"
                        pairs.append(f"{h}: {val}")
                    rows_text.append(", ".join(pairs))
                else:
                    rows_text.append(" | ".join(cells))
            if rows_text:
                parts.append(f"【工作表: {sheet_name}】\n" + "\n\n".join(rows_text))
        wb.close()
        full_text = "\n\n".join(parts)
        return ParsedDocument(file_name=file_name, total_pages=len(parts),
                            full_text=full_text, table_headers=headers)
    except Exception as exc:
        raise ValueError(f"无法解析 Excel: {file_name}") from exc


# ---------------------------------------------------------------------------
# Excel (.xls) —— 旧格式支持（xlrd）
# ---------------------------------------------------------------------------

def _parse_xls(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        import xlrd
    except ImportError:
        raise ValueError(
            f"不支持 .xls 文件: {file_name}。缺少 xlrd 库。"
        )
    try:
        wb = xlrd.open_workbook(file_contents=file_bytes)
        parts = []
        headers = {}
        for sheet_name in wb.sheet_names():
            ws = wb.sheet_by_name(sheet_name)
            if ws.nrows == 0:
                continue
            # 找表头行
            header_row = None
            header_cells = []
            for row_idx in range(ws.nrows):
                cells = [
                    str(ws.cell_value(row_idx, col_idx)).strip()
                    for col_idx in range(ws.ncols)
                    if ws.cell_value(row_idx, col_idx) != ""
                    and str(ws.cell_value(row_idx, col_idx)).strip()
                ]
                if len(cells) >= 2:
                    header_row = row_idx
                    header_cells = cells
                    break
            if header_row is not None:
                headers[sheet_name] = " | ".join(header_cells)
            rows_text = []
            for row_idx in range(ws.nrows):
                cells = [
                    str(ws.cell_value(row_idx, col_idx)).strip()
                    for col_idx in range(ws.ncols)
                    if ws.cell_value(row_idx, col_idx) != ""
                    and str(ws.cell_value(row_idx, col_idx)).strip()
                ]
                if not cells:
                    continue
                if row_idx == header_row:
                    rows_text.append(" | ".join(cells))
                elif header_cells and len(cells) <= len(header_cells):
                    pairs = []
                    for j, val in enumerate(cells):
                        h = header_cells[j] if j < len(header_cells) else f"列{j+1}"
                        pairs.append(f"{h}: {val}")
                    rows_text.append(", ".join(pairs))
                else:
                    rows_text.append(" | ".join(cells))
            if rows_text:
                parts.append(f"【工作表: {sheet_name}】\n" + "\n\n".join(rows_text))
        full_text = "\n\n".join(parts)
        return ParsedDocument(file_name=file_name, total_pages=len(parts),
                            full_text=full_text, table_headers=headers)
    except Exception as exc:
        raise ValueError(f"无法解析 Excel 文件: {file_name}") from exc


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------

def _parse_pptx(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        slides_text = []
        for idx, slide in enumerate(prs.slides, 1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            lines.append(t)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            lines.append(" | ".join(cells))
            if lines:
                slides_text.append(f"【幻灯片 {idx}】\n" + "\n".join(lines))
        full_text = "\n\n".join(slides_text)
        return ParsedDocument(file_name=file_name, total_pages=len(slides_text) or 1, full_text=full_text)
    except Exception as exc:
        raise ValueError(f"无法解析 PPT: {file_name}") from exc


# ---------------------------------------------------------------------------
# TXT / 文本类
# ---------------------------------------------------------------------------

def _parse_text(file_bytes: bytes, file_name: str) -> ParsedDocument:
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return ParsedDocument(file_name=file_name, total_pages=1, full_text=text.strip())
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError(f"无法解码文本文件: {file_name}")


# ---------------------------------------------------------------------------
# 图片 OCR（EasyOCR）
# ---------------------------------------------------------------------------

_ocr_reader = None


def _get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        import easyocr
        logger.info("首次加载 OCR 模型（约需 30 秒）…")
        _ocr_reader = easyocr.Reader(["ch_sim", "en"], gpu=False, verbose=False)
        logger.info("OCR 模型就绪")
    return _ocr_reader


def _parse_image(file_bytes: bytes, file_name: str) -> ParsedDocument:
    try:
        reader = _get_ocr_reader()
        import numpy as np
        img = Image.open(io.BytesIO(file_bytes))
        arr = np.array(img)
        results = reader.readtext(arr)
        lines = [text for (_bbox, text, _conf) in results if text.strip()]
        if lines:
            full_text = "\n".join(lines)
            logger.info("OCR 完成: %s → %d 行", file_name, len(lines))
            return ParsedDocument(file_name=file_name, total_pages=1, full_text=full_text)
        else:
            w, h = img.size
            return ParsedDocument(file_name=file_name, total_pages=1,
                                  full_text=f"[图片，未识别到文字] 文件名: {file_name}，尺寸: {w}x{h}")
    except Exception as exc:
        logger.warning("OCR 失败: %s — %s", file_name, exc)
        return ParsedDocument(file_name=file_name, total_pages=1,
                              full_text=f"[图片文件] {file_name}")
